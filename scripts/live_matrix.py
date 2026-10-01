"""Live test matrix runner: every registered Arelis tool singly, then chained.

Headless: builds the real seat (orchestrator + agent loop + production tool
registry + qwen3.5:9b through Ollama) and publishes USER_MESSAGE like the
window does, auto-clicking Allow on the bus. Everything runs in an isolated
data dir (outputs/test_matrix/data) with a synthetic profile, so the owner's
memory, contacts, mail/SMS credentials, calendar and browser profile are never
touched. Items live in scripts/live_matrix_items.py.

Hard exclusions: BANNED tools are removed from the registry AND watched on the
bus. If one is ever called the whole run aborts with exit code 3.

Usage:
  python -X utf8 scripts/live_matrix.py --kinds single,slash,regress
  python -X utf8 scripts/live_matrix.py --kinds chain
  python -X utf8 scripts/live_matrix.py --retry-failed
  python -X utf8 scripts/live_matrix.py --only S01_calculator_pct,R01_units_f_to_c
  python -X utf8 scripts/live_matrix.py --emit-matrix
"""

from __future__ import annotations

import argparse
import asyncio
import json
import os
import re
import sqlite3
import sys
import time
import traceback
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
MATRIX = ROOT / "outputs" / "test_matrix"
DATA = MATRIX / "data"
WORK = MATRIX / "work"
os.environ["ARELIS_DATA_DIR"] = str(DATA)
os.environ["ARELIS_SOLAR_GL"] = "0"
os.environ.pop("QT_QPA_PLATFORM", None)
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from live_matrix_items import BANNED, ITEMS, Item

EMAIL = re.compile(r"\b[\w.+-]+@[\w.-]+\.\w+\b")
PHONE = re.compile(r"(?:\+?1[\s.-]?)?\(?\d{3}\)?[\s.-]\d{3}[\s.-]\d{4}")
KIND_ORDER = {"single": 0, "regress": 0, "slash": 1, "chain": 2}


def snip(text: str, n: int = 300) -> str:
    one = " ".join((text or "").split())
    one = EMAIL.sub("[email]", one)
    one = PHONE.sub("[phone]", one)
    one = one.encode("ascii", "replace").decode("ascii")
    return one if len(one) <= n else one[: n - 3] + "..."


# ---------------------------------------------------------------- fixtures
def make_fixtures() -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    (DATA / "data").mkdir(parents=True, exist_ok=True)
    (DATA / "data" / "profile.yaml").write_text(
        "location:\n  city: Springfield\n  region: Illinois\n  country: US\n"
        "  latitude: 39.7817\n  longitude: -89.6501\n  timezone: America/Chicago\n",
        encoding="ascii",
    )
    (WORK / "sales.csv").write_text(
        "region,revenue\nnorth,100\nnorth,150\nsouth,250\neast,100\nwest,150\n", encoding="ascii"
    )
    (WORK / "hello_script.py").write_text("print(4242)\n", encoding="ascii")
    (MATRIX / "arelis-tasks.json").write_text(
        json.dumps({"tasks": {"hello": ["work/hello_script.py"]}}), encoding="ascii"
    )
    from fpdf import FPDF

    for name, code in (("a.pdf", "ALPHA-123"), ("b.pdf", "BRAVO-456")):
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=18)
        pdf.cell(0, 12, f"Matrix code: {code}")
        pdf.output(str(WORK / name))
    import docx

    d = docx.Document()
    d.add_paragraph("Project codename: FALCON-77.")
    d.add_paragraph("Approved budget: 4500 dollars.")
    d.save(str(WORK / "memo.docx"))
    from PIL import Image, ImageDraw, ImageFont

    def font(size: int):
        for cand in ("arial.ttf", "C:/Windows/Fonts/arial.ttf", "DejaVuSans.ttf"):
            try:
                return ImageFont.truetype(cand, size)
            except OSError:
                continue
        return ImageFont.load_default()

    img = Image.new("RGB", (1000, 300), "white")
    dr = ImageDraw.Draw(img)
    dr.text((40, 40), "INVOICE 88", fill="black", font=font(64))
    dr.text((40, 150), "Amount due: $1250.00", fill="black", font=font(64))
    img.save(WORK / "invoice.png")
    img = Image.new("RGB", (600, 400), "white")
    dr = ImageDraw.Draw(img)
    dr.ellipse((60, 80, 280, 300), fill=(220, 20, 20))
    dr.rectangle((330, 100, 540, 310), fill=(20, 40, 220))
    img.save(WORK / "shapes.png")


# ----------------------------------------------------------------- scoring
def fold(text: str) -> str:
    from arelis.eval.conversation import _fold_answer

    t = re.sub(r"(?<=\d),(?=\d{3})", "", text or "")
    return _fold_answer(t)


def seq_ok(called: list[str], expected: list[tuple]) -> bool:
    i = 0
    for name in called:
        if i < len(expected) and name in expected[i]:
            i += 1
    return i == len(expected)


def read_text_any(path: Path) -> str:
    suf = path.suffix.lower()
    try:
        if suf == ".pdf":
            from pypdf import PdfReader

            return " ".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
        if suf == ".docx":
            import docx

            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.extend(c.text for c in row.cells)
            return " ".join(parts)
        if suf == ".xlsx":
            import openpyxl

            wb = openpyxl.load_workbook(str(path), data_only=True)
            vals = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    vals.extend(str(v) for v in row if v is not None)
            return " ".join(vals)
        if suf == ".png":
            from PIL import Image

            with Image.open(path) as im:
                im.verify()
            return "PNG-OK"
        return path.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:
        return f"<<unreadable {type(exc).__name__}>>"


SKIP_DIRS = {"browser-profile", "__pycache__", "logs", "backups", "tool_cache", "earth"}


def snapshot() -> dict[str, tuple[int, int]]:
    snap: dict[str, tuple[int, int]] = {}
    for base, dirs, files in os.walk(MATRIX):
        rel = Path(base).relative_to(MATRIX).parts
        dirs[:] = [d for d in dirs if d not in SKIP_DIRS and not (rel == ("data",) and d == "data")]
        for f in files:
            p = Path(base) / f
            try:
                st = p.stat()
            except OSError:
                continue
            snap[str(p)] = (st.st_mtime_ns, st.st_size)
    return snap


def check_files(item: Item, before: dict, after: dict) -> tuple[list[str], list[str]]:
    changed = [Path(p) for p, sig in after.items() if before.get(p) != sig]
    reasons, found = [], []
    for suffix, needle in item.files:
        ok = False
        for p in changed:
            if suffix and p.suffix.lower() != suffix:
                continue
            if (
                p.suffix.lower()
                in {".db", ".wal", ".shm", ".yaml", ".json", ".jsonl", ".lock", ".pid", ".log"}
                and suffix == ""
            ):
                continue
            body = read_text_any(p)
            if needle.lower() in body.lower():
                ok = True
                found.append(str(p.relative_to(MATRIX)).replace("\\", "/"))
                break
        if not ok:
            reasons.append(f"no new/changed {suffix or 'any'} file containing {needle!r}")
    return reasons, found


def check_db(item: Item, db_path: Path) -> list[str]:
    reasons = []
    for table, col, needle in item.db:
        try:
            con = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
            rows = con.execute(f"select {col} from {table}").fetchall()
            con.close()
            if not any(needle.lower() in str(r[0]).lower() for r in rows):
                reasons.append(f"memory.db {table}.{col} has no row containing {needle!r}")
        except Exception as exc:
            reasons.append(f"memory.db check failed: {type(exc).__name__}")
    return reasons


# ------------------------------------------------------------------ runner
class Cap:
    def __init__(self) -> None:
        self.reset()

    def reset(self) -> None:
        self.tools: list[str] = []
        self.results: list[tuple[str, bool | None]] = []
        self.trace: list[str] = []
        self.final = ""
        self.error = ""
        self.done = asyncio.Event()
        self.confirms: list[str] = []
        self.timed_out = False


async def run_matrix(args) -> int:
    from arelis.config import load_config
    from arelis.core.bus import EventBus
    from arelis.core.events import Event, EventType
    from arelis.core.seat import build_seat
    from arelis.llm import prefix_warmup_for, run_model_preflight, run_model_warmup

    make_fixtures()
    cfg = load_config()
    cfg["workspace"] = {
        "roots": [{"name": "matrix", "path": str(MATRIX), "read_only": False}],
        "named_roots": [{"name": "matrix", "path": str(MATRIX), "read_only": False}],
    }
    tools_cfg = cfg.setdefault("tools", {})
    for key in (
        "email",
        "sms",
        "image",
        "calendar",
        "clipboard",
        "desktop",
        "briefing",
        "schedule",
    ):
        tools_cfg.setdefault(key, {})["enabled"] = False

    bus = EventBus()
    cap = Cap()
    state = {"violation": ""}

    async def on_start(ev: Event) -> None:
        name = str((ev.payload or {}).get("tool") or "")
        if not name:
            return
        cap.tools.append(name)
        if name in BANNED:
            state["violation"] = name
            await bus.publish(Event(EventType.TURN_CANCEL, {}))

    async def on_result(ev: Event) -> None:
        p = ev.payload or {}
        cap.results.append((str(p.get("tool") or ""), p.get("ok")))
        cap.trace.append(
            f"{p.get('tool')} ok={p.get('ok')} :: " + snip(str(p.get("output") or ""), 160)
        )

    async def on_done(ev: Event) -> None:
        cap.final = str((ev.payload or {}).get("text") or "")
        cap.done.set()

    async def on_err(ev: Event) -> None:
        if (ev.payload or {}).get("scope") == "voice":
            return
        cap.error = str((ev.payload or {}).get("message") or "error")
        cap.done.set()

    async def on_confirm(ev: Event) -> None:
        p = ev.payload or {}
        tool = str(p.get("tool") or "")
        cap.confirms.append(tool)
        banned = tool in BANNED or tool == "external_read"
        if tool in BANNED:
            state["violation"] = tool
        await bus.publish(
            Event(
                EventType.TOOL_CONFIRM_REPLY,
                {"id": p.get("id"), "decision": "skip" if banned else "allow", "allow_turn": False},
            )
        )

    bus.subscribe(EventType.TOOL_START, on_start)
    bus.subscribe(EventType.TOOL_RESULT, on_result)
    bus.subscribe(EventType.ASSISTANT_DONE, on_done)
    bus.subscribe(EventType.ERROR, on_err)
    bus.subscribe(EventType.TOOL_CONFIRM, on_confirm)

    seat = build_seat(cfg, profile="cli", bus=bus)
    removed = []
    for name in sorted(BANNED | {"image", "vision_camera"}):
        if seat.tools.get(name) is not None:
            seat.tools._tools.pop(name, None)
            removed.append(name)
    print("registered:", sorted(seat.tools.names()), flush=True)
    print("removed (banned/excluded):", removed, flush=True)
    bus_task = asyncio.create_task(bus.run())
    seat.router.arm_warmup()
    try:
        await run_model_preflight(bus, seat.router.provider, cfg.get("models"))
        await run_model_warmup(bus, seat.router, prefix=prefix_warmup_for(cfg, seat.tools))
    except Exception as exc:
        print("warmup warn:", type(exc).__name__, flush=True)
    finally:
        seat.router.mark_warmup_done()
    db_path = Path(seat.store.path)

    async def one_attempt(item: Item) -> dict:
        seat.memory.clear()
        try:
            seat.router.clear_sticky()
        except Exception:
            pass
        before = snapshot()
        all_tools: list[str] = []
        all_results: list[tuple[str, bool | None]] = []
        trace: list[str] = []
        finals: list[str] = []
        errors: list[str] = []
        timed_out = False
        t0 = time.perf_counter()
        for prompt in item.prompts:
            if prompt == "@clear":
                seat.memory.clear()
                continue
            cap.reset()
            await bus.publish(Event(EventType.USER_MESSAGE, {"text": prompt}))
            try:
                await asyncio.wait_for(cap.done.wait(), item.timeout)
            except TimeoutError:
                timed_out = True
                await bus.publish(Event(EventType.TURN_CANCEL, {}))
                try:
                    await asyncio.wait_for(cap.done.wait(), 90)
                except TimeoutError:
                    pass
            await asyncio.sleep(1.5)
            all_tools += cap.tools
            all_results += cap.results
            trace += cap.trace
            finals.append(cap.final)
            if cap.error:
                errors.append(cap.error)
            if state["violation"] or timed_out:
                break
        secs = round(time.perf_counter() - t0, 1)
        after = snapshot()
        final = finals[-1] if finals else ""
        reasons: list[str] = []
        if state["violation"]:
            reasons.append(f"BANNED tool touched: {state['violation']}")
        if timed_out:
            reasons.append(f"turn timeout after {item.timeout}s")
        if errors and not final:
            reasons.append("error: " + snip(errors[-1], 120))
        if item.tools:
            if not seq_ok(all_tools, item.tools):
                reasons.append(
                    f"expected tools {['|'.join(g) for g in item.tools]} in order, got {all_tools or '-'}"
                )
            else:
                for grp in item.tools:
                    recs = [ok for (n, ok) in all_results if n in grp]
                    if recs and not any(ok is not False for ok in recs):
                        reasons.append(f"{'|'.join(grp)} returned ok=False")
        elif not item.no_tools_ok and not all_tools:
            pass
        low = fold(final)
        for needle in item.ans_all:
            if fold(needle) not in low:
                reasons.append(f"answer missing {needle!r}")
        if item.ans_any and not any(fold(n) in low for n in item.ans_any):
            reasons.append(f"answer missing any of {list(item.ans_any)!r}")
        f_reasons, found = check_files(item, before, after)
        reasons += f_reasons
        reasons += check_db(item, db_path)
        return {
            "pass": not reasons,
            "reasons": reasons,
            "tools": all_tools,
            "final": snip(final, 320),
            "secs": secs,
            "files": found,
            "trace": trace,
            "confirms": list(cap.confirms),
        }

    out_json = MATRIX / "results.json"
    results: dict = {}
    if out_json.is_file():
        try:
            results = json.loads(out_json.read_text(encoding="ascii"))
        except Exception:
            results = {}
    items = sorted(ITEMS, key=lambda i: (KIND_ORDER[i.kind], i.timeout >= 900))
    if args.only:
        want = {x.strip() for x in args.only.split(",") if x.strip()}
        items = [i for i in items if i.id in want]
    if args.kinds:
        ks = {k.strip() for k in args.kinds.split(",")}
        items = [i for i in items if i.kind in ks]
    if args.retry_failed:
        items = [
            i
            for i in items
            if i.id in results
            and results[i.id]["status"] == "FAIL"
            and len(results[i.id]["attempts"]) < 2
        ]
    exit_code = 0
    for n, item in enumerate(items, 1):
        if item.skip:
            results[item.id] = {
                "id": item.id,
                "kind": item.kind,
                "tool": item.tool,
                "prompt": item.prompts[-1] if len(item.prompts) == 1 else " || ".join(item.prompts),
                "status": "SKIPPED",
                "reason": item.skip,
                "attempts": [],
                "tools": [],
                "final": "",
            }
            print(f"[{n}/{len(items)}] {item.id} SKIPPED: {item.skip}", flush=True)
            write_results(results)
            continue
        print(f"[{n}/{len(items)}] {item.id} ({item.tool}) ...", flush=True)
        try:
            att = await one_attempt(item)
        except Exception as exc:
            traceback.print_exc()
            att = {
                "pass": False,
                "reasons": [f"runner exception {type(exc).__name__}: {exc}"],
                "tools": [],
                "final": "",
                "secs": 0,
                "files": [],
            }
        rec = results.get(item.id) or {"attempts": []}
        rec.update(
            {
                "id": item.id,
                "kind": item.kind,
                "tool": item.tool,
                "prompt": " || ".join(item.prompts),
                "notes": item.notes,
            }
        )
        if args.retry_failed or "attempts" not in rec:
            pass
        rec["attempts"] = (rec.get("attempts") or []) + [att]
        atts = rec["attempts"]
        if att["pass"]:
            rec["status"] = "PASS"
            rec["flaky"] = len(atts) > 1 and not atts[0]["pass"]
            rec["reason"] = "flaky: passed on rerun" if rec["flaky"] else "ok"
        else:
            rec["status"] = "FAIL"
            rec["reason"] = "; ".join(att["reasons"])
        rec["tools"] = att["tools"]
        rec["final"] = att["final"]
        results[item.id] = rec
        print(
            f"    {rec['status']} {att['secs']}s tools={att['tools']} :: {snip(rec['reason'], 160)}",
            flush=True,
        )
        write_results(results)
        if state["violation"]:
            print(
                f"HARD FAIL: banned tool {state['violation']} was called. Aborting run.", flush=True
            )
            exit_code = 3
            break
    write_results(results)
    bus.stop()
    bus_task.cancel()
    try:
        await seat.router.close()
    except Exception:
        pass
    return exit_code


def write_results(results: dict) -> None:
    MATRIX.mkdir(parents=True, exist_ok=True)
    (MATRIX / "results.json").write_text(json.dumps(results, indent=1), encoding="ascii")
    order = {i.id: n for n, i in enumerate(ITEMS)}
    rows = sorted(results.values(), key=lambda r: order.get(r["id"], 999))
    p = sum(1 for r in rows if r["status"] == "PASS")
    f = sum(1 for r in rows if r["status"] == "FAIL")
    s = sum(1 for r in rows if r["status"] == "SKIPPED")
    lines = [
        f"# Arelis live test matrix results  {datetime.now().strftime('%Y-%m-%d %H:%M')} ET",
        "",
        f"{p} PASS / {f} FAIL / {s} SKIPPED / {len(rows)} run (of {len(ITEMS)} defined)",
        "",
        "| id | prompt | tools in order | result | reason | answer snippet |",
        "|----|--------|----------------|--------|--------|----------------|",
    ]
    for r in rows:
        prompt = snip(r.get("prompt", ""), 110).replace("|", "/")
        tools = ", ".join(r.get("tools") or []) or "-"
        reason = snip(r.get("reason", ""), 140).replace("|", "/")
        fin = snip(r.get("final", ""), 120).replace("|", "/")
        flag = r["status"] + (" (flaky)" if r.get("flaky") else "")
        lines.append(f"| {r['id']} | {prompt} | {tools} | {flag} | {reason} | {fin} |")
    (MATRIX / "results.md").write_text("\n".join(lines) + "\n", encoding="ascii")


def emit_matrix() -> None:
    lines = [
        "# Arelis live test matrix",
        "",
        "Runner: `scripts/live_matrix.py` (items in `scripts/live_matrix_items.py`).",
        "Workspace root for the run: `outputs/test_matrix` (fixtures in `work/`). Isolated data dir "
        "`outputs/test_matrix/data` with a synthetic Springfield profile.",
        "Pass = expected tools appear in order (alternatives in parentheses) AND answer substring(s) "
        "AND/OR new artifact content AND/OR memory.db row are present.",
        "",
    ]
    for kind, title in (
        ("single", "Single-tool"),
        ("regress", "Regression"),
        ("slash", "Slash commands / features"),
        ("chain", "Chained (2-5 tools)"),
    ):
        lines += [
            f"## {title}",
            "",
            "| id | prompt | expected tools (in order) | expected answer / artifact |",
            "|---|---|---|---|",
        ]
        for it in ITEMS:
            if it.kind != kind:
                continue
            tools = " > ".join("|".join(g) for g in it.tools) or "(none, handled by orchestrator)"
            exp = []
            if it.ans_all:
                exp.append("answer has " + " & ".join(repr(x) for x in it.ans_all))
            if it.ans_any:
                exp.append("answer has any of " + ", ".join(repr(x) for x in it.ans_any))
            for suf, nd in it.files:
                exp.append(f"new {suf or 'any'} file" + (f" containing {nd!r}" if nd else ""))
            for t, c, nd in it.db:
                exp.append(f"memory.db {t}.{c} contains {nd!r}")
            if it.skip:
                exp.append("SKIPPED: " + it.skip)
            prompt = " || ".join(it.prompts).replace("|", "/")
            lines.append(f"| {it.id} | {prompt} | {tools} | {'; '.join(exp) or 'tool called'} |")
        lines.append("")
    (MATRIX / "matrix.md").write_text("\n".join(lines), encoding="ascii", errors="replace")
    print("wrote", MATRIX / "matrix.md")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", default="")
    ap.add_argument("--kinds", default="")
    ap.add_argument("--retry-failed", action="store_true")
    ap.add_argument("--emit-matrix", action="store_true")
    args = ap.parse_args()
    if args.emit_matrix:
        emit_matrix()
        return 0
    return asyncio.run(run_matrix(args))


if __name__ == "__main__":
    raise SystemExit(main())
