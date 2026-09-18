"""Read text out of local Office files. No cloud. No python-pptx."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

_SLIDE_NAME = re.compile(r"ppt/slides/slide(\d+)\.xml$")
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def sniff_office_kind(path: Path) -> str:
    """Return 'docx' or 'pptx' from the zip, or '' if it is not Office.

    Filename is ignored. A Word file saved as notes.pdf is still a docx.
    """
    try:
        with zipfile.ZipFile(path) as zf:
            names = zf.namelist()
    except (OSError, zipfile.BadZipFile):
        return ""
    if "word/document.xml" in names:
        return "docx"
    if any(_SLIDE_NAME.match(name) for name in names):
        return "pptx"
    types = ""
    if "[Content_Types].xml" in names:
        try:
            with zipfile.ZipFile(path) as zf:
                types = zf.read("[Content_Types].xml").decode("utf-8", errors="replace")
        except (OSError, zipfile.BadZipFile, KeyError):
            types = ""
    if "wordprocessingml" in types:
        return "docx"
    if "presentationml" in types:
        return "pptx"
    return ""


def extract_docx_text(path: Path) -> str:
    """Paragraphs plus tables as readable cell text, not a binary blob."""
    from docx import Document

    doc = Document(str(path))
    chunks: list[str] = []
    for paragraph in doc.paragraphs:
        text = (paragraph.text or "").strip()
        if text:
            chunks.append(text)
    for table in doc.tables:
        rendered = _render_docx_table(table)
        if rendered:
            chunks.append(rendered)
    return "\n\n".join(chunks).strip()


def extract_pptx_slides(path: Path) -> list[str]:
    """One string per slide, ordered by slide number. Empty slides stay ''."""
    slides: dict[int, str] = {}
    try:
        with zipfile.ZipFile(path) as zf:
            for name in zf.namelist():
                match = _SLIDE_NAME.match(name)
                if not match:
                    continue
                number = int(match.group(1))
                slides[number] = _slide_text(zf.read(name))
    except (OSError, zipfile.BadZipFile) as exc:
        raise ValueError(f"not a pptx zip: {exc}") from exc
    if not slides:
        return []
    last = max(slides)
    return [slides.get(idx, "") for idx in range(1, last + 1)]


def _render_docx_table(table: object) -> str:
    rows_out: list[str] = []
    for row in getattr(table, "rows", []):
        cells = []
        for cell in getattr(row, "cells", []):
            text = " ".join((cell.text or "").split())
            cells.append(text)
        if any(cells):
            rows_out.append(" | ".join(cells))
    if not rows_out:
        return ""
    return "[table]\n" + "\n".join(rows_out)


def _slide_text(xml_bytes: bytes) -> str:
    try:
        root = ET.fromstring(xml_bytes)
    except ET.ParseError:
        return ""
    parts: list[str] = []
    for node in root.iter(f"{{{_A_NS}}}t"):
        text = (node.text or "").strip()
        if text:
            parts.append(text)
    return "\n".join(parts).strip()
