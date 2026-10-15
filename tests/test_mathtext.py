"""Unicode math is the one form chat, PDF, Word, CSV, and markdown share."""

from __future__ import annotations

from pathlib import Path

import pytest

from arelis.mathtext import flatten_latex
from arelis.tools.document import DocumentTool
from arelis.ui.markdown import render_markdown
from arelis.workspace import WorkspaceRoots


def test_boxed_does_not_print_the_command_name() -> None:
    out = flatten_latex(r"$$\boxed{\theta_{opt} \approx 35.3^\circ}$$")
    assert "boxed" not in out
    assert "θ" in out
    assert "≈" in out
    assert "35.3" in out
    assert "°" in out
    assert r"\boxed" not in out


def test_integral_delimiters_leave() -> None:
    out = flatten_latex(
        r"The integral of \( x^2 \) is \[\int x^2 \, dx = \frac{x^3}{3} + C\]"
    )
    assert r"\(" not in out
    assert r"\[" not in out
    assert "x²" in out
    assert "∫" in out
    assert "x³" in out
    assert "/3" in out


def test_windows_path_is_not_eaten() -> None:
    out = flatten_latex(r"read C:\Users\you\notes.md then \frac{1}{2}")
    assert r"\Users" in out
    assert "1/2" in out
    assert r"\frac" not in out
    assert "C:\\input" in flatten_latex(r"open C:\input next")
    assert "∈" not in flatten_latex(r"open C:\input next")
    assert r"folder\log" in flatten_latex(r"see folder\log then done")


def test_chat_centers_display_math() -> None:
    html = render_markdown(r"The energy is $$\frac{1}{2}mv^2$$ as usual.")
    assert "text-align:center" in html
    assert "font-style:italic" in html
    assert "½" in html or "1/2" in html
    assert "mv²" in html or "mv^2" in html
    assert r"\frac" not in html
    assert "$$" not in html


@pytest.fixture
def doc_tool(tmp_path, monkeypatch):
    root = tmp_path / "project"
    root.mkdir()
    data = tmp_path / "appdata"
    data.mkdir()
    monkeypatch.setenv("ARELIS_DATA_DIR", str(data))
    workspace = WorkspaceRoots.from_config(
        {"workspace": {"roots": [{"name": "project", "path": str(root)}]}}
    )
    return DocumentTool(workspace)


@pytest.mark.asyncio
async def test_pdf_word_csv_md_all_flatten_tex(doc_tool) -> None:
    body = (
        "The closed form is\n\n"
        r"$$\frac{x^{3}}{6} + 25x\log(x-3)$$" + "\n\n"
        "and the root is " + r"$\sqrt{x+1}$."
    )

    md = await doc_tool.run(format="md", title="Math", body=body)
    assert md.ok, md.output
    md_text = Path(md.data["abs_path"]).read_text(encoding="utf-8")
    assert r"\frac" not in md_text
    assert "log" in md_text
    assert "x³" in md_text
    assert "√(x+1)" in md_text

    txt = await doc_tool.run(format="txt", title="Math txt", body=body)
    assert txt.ok, txt.output
    txt_text = Path(txt.data["abs_path"]).read_text(encoding="utf-8")
    assert r"\frac" not in txt_text
    assert "√(x+1)" in txt_text

    pdf = await doc_tool.run(format="pdf", title="Math pdf", body=body)
    assert pdf.ok, pdf.output
    from pypdf import PdfReader

    pages = PdfReader(str(Path(pdf.data["abs_path"]))).pages
    pdf_text = "\n".join(page.extract_text() or "" for page in pages)
    assert r"\frac" not in pdf_text
    assert "log" in pdf_text

    docx = await doc_tool.run(format="docx", title="Math docx", body=body)
    assert docx.ok, docx.output
    from docx import Document

    word = Document(str(Path(docx.data["abs_path"])))
    word_text = "\n".join(p.text for p in word.paragraphs)
    assert r"\frac" not in word_text
    assert "log" in word_text
    assert any(run.italic for p in word.paragraphs for run in p.runs)

    csv = await doc_tool.run(
        format="csv",
        title="Math csv",
        rows='[["formula"], ["$$\\\\frac{a}{b}$$"]]',
    )
    assert csv.ok, csv.output
    csv_text = Path(csv.data["abs_path"]).read_text(encoding="utf-8-sig")
    assert r"\frac" not in csv_text
    assert "a/b" in csv_text

    xlsx = await doc_tool.run(
        format="xlsx",
        title="Math xlsx",
        rows='[["formula"], ["$$\\\\sqrt{x+1}$$"]]',
    )
    assert xlsx.ok, xlsx.output
    from openpyxl import load_workbook

    book = load_workbook(Path(xlsx.data["abs_path"]))
    assert book.active["A2"].value == "√(x+1)"

    source = Path(md.data["abs_path"])
    exported = await doc_tool.run(
        format="pdf",
        title="From md",
        from_path=str(source),
    )
    assert exported.ok, exported.output
    from pypdf import PdfReader

    from_text = "\n".join(
        page.extract_text() or ""
        for page in PdfReader(str(Path(exported.data["abs_path"]))).pages
    )
    assert r"\frac" not in from_text
    assert "log" in from_text

    table_body = "| a | b |\n| --- | --- |\n| $\\frac{1}{2}$ | ok |\n"
    table_doc = await doc_tool.run(format="docx", title="Table math", body=table_body)
    assert table_doc.ok, table_doc.output
    table_word = Document(str(Path(table_doc.data["abs_path"])))
    cells = [[c.text for c in row.cells] for t in table_word.tables for row in t.rows]
    assert ["1/2", "ok"] in cells
