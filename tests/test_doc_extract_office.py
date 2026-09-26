"""Office + form-field coverage for doc_extract (roadmap 4.11)."""

from __future__ import annotations

import zipfile
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape as xml_escape

import pytest
from PIL import Image

from arelis.tools.doc_extract import DocExtractTool, build_simple_pdf_bytes
from arelis.workspace import ResolvedPath


def _write_docx(
    path: Path,
    paragraph: str = "",
    table: list[list[str]] | None = None,
) -> Path:
    from docx import Document

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    if paragraph:
        doc.add_paragraph(paragraph)
    if table:
        width = max(len(row) for row in table)
        grid = doc.add_table(rows=len(table), cols=width)
        for r_i, row in enumerate(table):
            for c_i in range(width):
                grid.cell(r_i, c_i).text = row[c_i] if c_i < len(row) else ""
    doc.save(str(path))
    return path


def _write_pptx(path: Path, slides: list[str]) -> Path:
    """Tiny real zip+xml deck. No python-pptx."""
    path.parent.mkdir(parents=True, exist_ok=True)
    overrides = []
    for idx, _text in enumerate(slides, start=1):
        overrides.append(
            f'<Override PartName="/ppt/slides/slide{idx}.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.'
            'presentationml.slide+xml"/>'
        )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/ppt/presentation.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.'
        'presentationml.presentation.main+xml"/>'
        + "".join(overrides)
        + "</Types>"
    )
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr(
            "ppt/presentation.xml",
            '<?xml version="1.0"?><p:presentation '
            'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>',
        )
        for idx, text in enumerate(slides, start=1):
            zf.writestr(f"ppt/slides/slide{idx}.xml", _slide_xml(text))
    return path


def _slide_xml(text: str) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        "<p:cSld><p:spTree><p:sp><p:txBody>"
        f"<a:p><a:r><a:t>{xml_escape(text)}</a:t></a:r></a:p>"
        "</p:txBody></p:sp></p:spTree></p:cSld></p:sld>"
    )


def _write_form_pdf(path: Path) -> Path:
    content = b"BT /F1 12 Tf 72 720 Td (A form) Tj ET"
    parts = [
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R /AcroForm 6 0 R >>\nendobj\n",
        b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n",
        (
            b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Annots [7 0 R] /Resources << /Font << /F1 5 0 R >> "
            b">> >>\nendobj\n"
        ),
        (
            f"4 0 obj\n<< /Length {len(content)} >>\nstream\n".encode("latin-1")
            + content
            + b"\nendstream\nendobj\n"
        ),
        b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n",
        b"6 0 obj\n<< /Fields [7 0 R] /NeedAppearances true >>\nendobj\n",
        (
            b"7 0 obj\n<< /Type /Annot /Subtype /Widget /FT /Tx /T (FullName) "
            b"/V (Alice Example) /Ff 0 /Rect [72 680 300 700] /P 3 0 R /F 4 "
            b"/DA (/Helv 12 Tf 0 g) >>\nendobj\n"
        ),
    ]
    header = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n"
    body = b"".join(parts)
    offsets = [0]
    pos = len(header)
    for obj in parts:
        offsets.append(pos)
        pos += len(obj)
    xref = [f"xref\n0 {len(offsets)}\n".encode("latin-1"), b"0000000000 65535 f \n"]
    for off in offsets[1:]:
        xref.append(f"{off:010d} 00000 n \n".encode("latin-1"))
    trailer = (
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
        f"startxref\n{pos}\n%%EOF\n"
    ).encode("latin-1")
    path.write_bytes(header + body + b"".join(xref) + trailer)
    return path


def _png(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (8, 8), "white")
    buf = BytesIO()
    image.save(buf, format="PNG")
    path.write_bytes(buf.getvalue())
    return path


def _join_only(path_str: str) -> ResolvedPath:
    """Mutant (1): skip WorkspaceRoots.resolve_read."""
    path = Path(path_str)
    return ResolvedPath(path=path, root_name="leaked", root=path.parent)


@pytest.mark.asyncio
async def test_docx_paragraph_text(tmp_path: Path) -> None:
    dest = _write_docx(tmp_path / "notes.docx", "hello from word")
    tool = DocExtractTool([str(tmp_path)])
    result = await tool.run(path=str(dest))
    assert result.ok
    assert result.data["source"] == "docx"
    assert "hello from word" in result.output


@pytest.mark.asyncio
async def test_docx_table_cells_are_readable(tmp_path: Path) -> None:
    dest = _write_docx(
        tmp_path / "grid.docx",
        table=[["alpha", "beta"], ["gamma", "delta"]],
    )
    tool = DocExtractTool([str(tmp_path)])
    result = await tool.run(path=str(dest))
    assert result.ok
    assert result.data["source"] == "docx"
    for cell in ("alpha", "beta", "gamma", "delta"):
        assert cell in result.output
    assert "[table]" in result.output


@pytest.mark.asyncio
async def test_docx_outside_workspace_is_refused(tmp_path: Path) -> None:
    inside = tmp_path / "proj"
    outside = tmp_path / "elsewhere"
    inside.mkdir()
    secret = _write_docx(outside / "secret.docx", "classified payload")
    tool = DocExtractTool([str(inside)])
    result = await tool.run(path=str(secret))
    assert not result.ok
    assert "outside allowed workspace roots" in result.output.lower()
    assert "classified payload" not in result.output


@pytest.mark.asyncio
async def test_plain_path_join_resolve_breaks_containment(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Mutant (1): Path join instead of resolve_read. Containment dies."""
    inside = tmp_path / "proj"
    outside = tmp_path / "elsewhere"
    inside.mkdir()
    secret = _write_docx(outside / "secret.docx", "classified payload")
    tool = DocExtractTool([str(inside)])
    honest = await tool.run(path=str(secret))
    assert not honest.ok

    monkeypatch.setattr(tool, "_resolve", _join_only)
    leaked = await tool.run(path=str(secret))
    assert leaked.ok
    assert "classified payload" in leaked.output


@pytest.mark.asyncio
async def test_docx_named_pdf_is_not_parsed_as_pdf(tmp_path: Path) -> None:
    """Pin the lie: content wins. notes.pdf that is a docx is still a docx."""
    real = _write_docx(tmp_path / "real.docx", "secret paragraph")
    lie = tmp_path / "notes.pdf"
    lie.write_bytes(real.read_bytes())
    tool = DocExtractTool([str(tmp_path)])
    result = await tool.run(path=str(lie))
    assert result.ok
    assert result.data["source"] == "docx"
    assert "secret paragraph" in result.output
    assert result.data.get("source") != "text"
    assert result.data.get("source") != "ink"
    assert not result.data.get("page_images")


@pytest.mark.asyncio
async def test_docx_is_not_routed_through_pdf(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def boom(*_args: object, **_kwargs: object) -> None:
        raise AssertionError("docx routed through PDF extractor")

    monkeypatch.setattr(DocExtractTool, "_extract_pdf", boom)
    monkeypatch.setattr(DocExtractTool, "_extract_ink", boom)
    dest = _write_docx(tmp_path / "safe.docx", "stay off the pdf path")
    tool = DocExtractTool([str(tmp_path)])
    result = await tool.run(path=str(dest))
    assert result.ok
    assert result.data["source"] == "docx"
    assert "stay off the pdf path" in result.output


@pytest.mark.asyncio
async def test_pptx_slide_two_is_distinct(tmp_path: Path) -> None:
    dest = _write_pptx(tmp_path / "deck.pptx", ["agenda items", "budget numbers"])
    tool = DocExtractTool([str(tmp_path)])
    both = await tool.run(path=str(dest))
    assert both.ok
    assert both.data["source"] == "pptx"
    assert "slide 1:" in both.output
    assert "agenda items" in both.output
    assert "slide 2:" in both.output
    assert "budget numbers" in both.output

    only_two = await tool.run(path=str(dest), page_start=2, page_end=2)
    assert only_two.ok
    assert "budget numbers" in only_two.output
    assert "agenda items" not in only_two.output
    assert only_two.data["pages"] == [2]


@pytest.mark.asyncio
async def test_unsupported_exe_and_png_are_fail_tags(tmp_path: Path) -> None:
    exe = tmp_path / "setup.exe"
    exe.write_bytes(b"MZ\x90\x00not-a-document")
    png = _png(tmp_path / "shot.png")
    tool = DocExtractTool([str(tmp_path)])

    exe_result = await tool.run(path=str(exe))
    assert not exe_result.ok
    assert exe_result.data["fail_class"] == "fail:unsupported"
    assert exe_result.output.startswith("[fail:unsupported]")

    png_result = await tool.run(path=str(png))
    assert not png_result.ok
    assert png_result.data["fail_class"] == "fail:unsupported"
    assert png_result.output.startswith("[fail:unsupported]")
    assert "vision" in png_result.output.lower()


@pytest.mark.asyncio
async def test_pdf_form_fields_are_listed(tmp_path: Path) -> None:
    dest = _write_form_pdf(tmp_path / "form.pdf")
    tool = DocExtractTool([str(tmp_path)])
    result = await tool.run(path=str(dest))
    assert result.ok
    assert result.data["source"] == "text"
    assert result.data.get("form_fields", {}).get("FullName") == "Alice Example"
    assert "Alice Example" in result.output
    assert "form fields:" in result.output


@pytest.mark.asyncio
async def test_text_pdf_still_extracts_through_dispatch(tmp_path: Path) -> None:
    pdf = tmp_path / "note.pdf"
    pdf.write_bytes(build_simple_pdf_bytes("hello from the text layer"))
    tool = DocExtractTool([str(tmp_path)])
    result = await tool.run(path=str(pdf))
    assert result.ok
    assert result.data["source"] == "text"
    assert "hello from the text layer" in result.output
